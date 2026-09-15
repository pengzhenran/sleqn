#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_wechat_article.py —— 生成公众号简介文章的 Word 版
============================================================================

面向公众号「地球重力与人类生活（TVGG）」的读者，介绍《海平面指纹（SLF）
计算程序》，内容与 docs/使用说明.html 保持一致（数字都从那里核对过）。

    <python> tools\\make_wechat_article.py

产物：公众号文章_海平面指纹计算程序.docx（放在项目根目录）

配图取自 docs/使用说明_img/。文中「下载方式」一段留了占位，发布前改成实际
获取途径即可。
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, 'docs', '使用说明_img')
QR = os.path.join(ROOT, 'assets', '地球重力与人类生活TVGG.jpg')
OUT = os.path.join(ROOT, '公众号文章_海平面指纹计算程序.docx')

FONT = '微软雅黑'
DARK = RGBColor(0x22, 0x22, 0x33)
BLUE = RGBColor(0x1B, 0x3A, 0x5C)
GREY = RGBColor(0x66, 0x66, 0x77)
ACCENT = RGBColor(0x2F, 0x6F, 0xA8)


# ---------------------------------------------------------------- 排版小工具
# OOXML 的 CT_PPr / CT_RPr 对子元素顺序有 schema 规定（pBdr、shd 要排在
# spacing/ind/jc 之前，rFonts 要排在 b/sz/color 之前）。自己 makeelement + append
# 很容易插错位置，所以能用 python-docx 的顺序感知接口就用它
# （如 rPr.get_or_add_rFonts()）；python-docx 没提供 accessor 的（pBdr、shd）
# 就按下面的顺序表自己插。
_PPR_SEQ = (
    'w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
    'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd',
    'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
    'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
    'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
    'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
    'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
    'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange',
)


def _ppr_insert(pPr, el, tag):
    """按 CT_PPr 的顺序把 el 插到 pPr 里（找不到后继元素就追加到末尾）。"""
    pPr.insert_element_before(el, *_PPR_SEQ[_PPR_SEQ.index(tag) + 1:])


def _cn(run, name=FONT):
    """python-docx 只管西文字体，中文要单独设 w:eastAsia。

    用 ``run.font.name`` 与 ``rPr.get_or_add_rFonts()`` 这两个顺序感知的接口，
    不要自己 makeelement + append（CT_RPr 对子元素顺序有要求）。
    """
    run.font.name = name                       # 建/定位 w:rFonts（含 ascii、hAnsi）
    rPr = run._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn('w:eastAsia'), name)


def para(doc, text='', size=10.5, bold=False, color=DARK, align=None,
         space_after=8, space_before=0, indent=False, line=1.6, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        _cn(r)
    return p


def rich(doc, parts, size=10.5, align=None, space_after=8, line=1.6):
    """一段里混排「普通 + 加粗 + 强调色」。parts = [(文本, 样式), ...]"""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    for text, style in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = 'b' in style
        r.font.color.rgb = ACCENT if 'a' in style else (
            GREY if 'g' in style else DARK)
        _cn(r)
    return p


def h1(doc, text, sub=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(19)
    r.bold = True
    r.font.color.rgb = BLUE
    _cn(r)
    if sub:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        r = p.add_run(sub)
        r.font.size = Pt(10)
        r.font.color.rgb = GREY
        _cn(r)


def h2(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16)
    pf.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(13.5)
    r.bold = True
    r.font.color.rgb = BLUE
    _cn(r)
    # 段落下加一条细线，接近公众号小标题的观感
    pPr = p._element.get_or_add_pPr()
    bdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6',
        qn('w:space'): '4', qn('w:color'): 'B9CDE0'})
    bdr.append(bottom)
    _ppr_insert(pPr, bdr, 'w:pBdr')


def bullet(doc, text, lead=None):
    p = doc.add_paragraph(style='List Bullet')
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.5
    if lead:
        r = p.add_run(lead)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
        _cn(r)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK
    _cn(r)
    return p


def image(doc, filename, caption, width=15.0, path=None):
    src = path or os.path.join(IMG, filename)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(src, width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(14)
    r = cap.add_run(caption)
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    _cn(r)


def quote(doc, text):
    """导读框：左缩进 + 浅底纹。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5)
    pf.right_indent = Cm(0.5)
    pf.space_before = Pt(6)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.55
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x33, 0x44, 0x55)
    _cn(r)
    pPr = p._element.get_or_add_pPr()
    shd = pPr.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'F0F7F2'})
    _ppr_insert(pPr, shd, 'w:shd')
    return p


def table(doc, rows, widths=(4.6, 4.2, 3.4, 3.0)):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, txt in enumerate(row):
            cells[j].width = Cm(widths[j])
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(txt)
            r.font.size = Pt(9.5)
            r.bold = (i == 0)
            r.font.color.rgb = BLUE if i == 0 else DARK
            _cn(r)
    return t


# ---------------------------------------------------------------- 正文
def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)   # A4
    sec.left_margin = sec.right_margin = Cm(2.8)           # 正文宽 15.4 cm
    sec.top_margin = sec.bottom_margin = Cm(2.5)

    h1(doc, '陆地水一去不回，海平面却不是均匀上升',
       '「海平面指纹」计算程序 v1.0 发布 · 免费 · 离线可用')

    quote(doc,
          '假设格陵兰冰盖融掉一千亿吨水，全部流进海洋。全球海平面会上升多少？'
          '直觉答案是「总水量 ÷ 海洋面积」——一个到处都一样的数。\n'
          '但真实情况并不是这样：冰盖附近的海面可能反而下降，'
          '而一万公里外的热带海面上升得更多。这就是「海平面指纹」。')

    # ---------------- 1
    h2(doc, '一、什么是海平面指纹')
    rich(doc, [
        ('地球上的水在陆地、海洋、大气之间来回搬家时，海水并不会"均匀地"升降。'
         '有两个机制同时在起作用：', ''),
    ])
    bullet(doc, '质量增加的地方引力变强，会把附近的海水吸过去；冰盖融化则相反，'
                '引力减弱，海水被"松开"。', lead='① 自吸引　')
    bullet(doc, '地表的重量变了，固体地球会被压下去或弹回来，海床跟着升降。',
           lead='② 固体地球形变　')
    rich(doc, [
        ('两者叠加，就形成了"近处与远处反号"的空间图案——', ''),
        ('这就是海平面指纹。', 'b'),
    ])
    image(doc, 'fig_principle.png',
          '图 1　海平面方程各物理量的含义：负荷把海面"吸"起来、'
          '把固体地球压下去，二者之差决定海面高度')

    rich(doc, [
        ('要把这个图案算出来，得解一个叫', ''),
        ('海平面方程', 'b'),
        ('的积分方程：', ''),
    ])
    para(doc, 'S(θ, ψ, t) = C(θ, ψ, t) · [ G^L(θ, ψ, t) − R^L(θ, ψ, t) ] + c(t)',
         size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    rich(doc, [
        ('其中 C 是海洋函数（海洋为 1、陆地为 0），G^L 是负荷引起的大地水准面异常，'
         'R^L 是固体地球表面的径向位移，c 是维持海水质量守恒的常数项——'
         '它需要在迭代中逐步修正。', ''),
    ], space_after=6)
    rich(doc, [
        ('换成大白话：', ''),
        ('已知地表哪里多了水、哪里少了水，求全球海面怎么变。', 'b'),
    ])

    # ---------------- 2
    h2(doc, '二、为什么值得算一算')
    bullet(doc, '卫星测高测到的是"总的海平面变化"，里面混着指纹效应；'
                '想从里面分离出真正的水量变化，就得把指纹算出来。',
           lead='分离信号　')
    bullet(doc, '做冰盖质量平衡、区域海平面预估、GRACE 数据的解释，'
                '都绕不开这一步。', lead='课题需要　')
    bullet(doc, '方程本身不难写，但要正确处理球谐展开、海洋函数、自吸引与负荷迭代、'
                '以及质量守恒，自己从头实现容易在细节上出错。', lead='自己写麻烦　')

    # ---------------- 3
    h2(doc, '三、这个工具做了什么')
    rich(doc, [
        ('把上面那套流程做成了一个', ''),
        ('带图形界面的 Windows 程序', 'b'),
        ('：准备三份文本文件（海陆掩膜、负荷勒夫数、质量变化网格），'
         '点一下「开始计算」，就得到全球海平面变化网格。', ''),
    ])
    image(doc, 'gui_01_overview.png',
          '图 2　程序主界面：左边填文件和参数，右边四个页签看结果')

    rich(doc, [
        ('输入的质量变化可以是 ', ''),
        ('GRACE/GRACE-FO 反演的陆地水储量变化', 'b'),
        ('，也可以是冰盖模型、地下水模型给出的任何等效水高网格；'
         '输出是每个 1° 格点上的海平面变化（厘米）。', ''),
    ])
    image(doc, 'fig_example.png',
          '图 3　一个合成算例：北极附近一块 5°×5° 的正载荷（相当于 65.6 Gt '
          '的质量增加）。载荷附近海面抬升，全球远场海面普遍下降')

    h2(doc, '四、几个用起来舒服的地方')
    bullet(doc, '装完双击就能用，不需要装 Python、不需要联网，'
                '运行过程不联网、不上传任何数据。', lead='开箱即用　')
    bullet(doc, '随机附带真实的 CSR GRACE/GRACE-FO RL06.3 mascon 示例数据'
                '（1° 网格、64800 个格点）与 PREM 负荷勒夫数，'
                '点一下「载入演示数据」就能把整条流程走一遍，'
                '全球算完只要十几秒。', lead='自带真实算例　')
    bullet(doc, '地图、经/纬向剖面、质量守恒自检、结果预览都在界面里；'
                '图片可直接导出 PNG/PDF，结果文件就是三列文本，'
                'MATLAB / Python / Surfer 都能读。', lead='结果看得见　')
    bullet(doc, '按 F1 直接打开中文使用说明（含全部界面截图、参数含义、'
                '常见问题），窗口内就能看，不用另开浏览器。', lead='中文说明　')
    bullet(doc, '可以自己指定球谐截断阶数（默认 180），'
                '换用不同阶数的勒夫数表；算例大小、网格间隔都是可调的。',
           lead='参数可调　')
    bullet(doc, '极移（自转）反馈可以一键开关：默认按 Kendall et al. (2005) 计入'
                '自转轴漂移引起的离心位扰动，与 Fortran 版和 gravity-toolkit 一致；'
                '取消勾选即得不含该反馈的经典解，方便做敏感性分析。',
           lead='极移可控　')

    # ---------------- 4
    h2(doc, '五、算得准不准？我们拿三套独立程序对过')
    rich(doc, [
        ('这是最该说清楚的一件事。我们把本程序算出的', ''),
        ('载荷球谐系数原样取出', 'b'),
        ('，分别喂给两套完全独立的第三方实现，再另与一个解析闭式解对照——'
         '这样两边用的输入逐位相同，差异只可能来自求解器本身。', ''),
    ])
    table(doc, [
        ['对照对象', '算例', '最大偏差（占峰值）', '相关系数'],
        ['解析闭式解', '合成算例', '0.029 %', '0.99999997'],
        ['gravity-toolkit', '合成算例', '0.030 %', '0.99999997'],
        ['gravity-toolkit', '真实 GRACE', '0.022 %', '0.99999999'],
        ['pyslfp 2.0.6', '全球海洋 + 光滑载荷', '0.537 %', '0.99999919'],
    ])
    rich(doc, [
        ('两个互相独立的参考之间只差 0.0045 %，说明它们确实是各自独立的正确实现；'
         '本程序与它们相差 ', ''),
        ('万分之二到万分之三', 'b'),
        ('。与 pyslfp 的 0.5 % 主要来自两边用的负荷勒夫数表不同'
         '（pyslfp 自带 PREM_4096，本程序用 Wang et al. 2012 PREM）'
         '与地球半径取值差异，两者勒夫数响应因子本身就差 0.19 %。', ''),
    ], space_after=10)
    image(doc, 'fig_verify.png',
          '图 4　三路对照：本程序（左上）与 gravity-toolkit（中上）算出的指纹'
          '几乎重合，差值（右上）只有峰值的 0.03 %')

    quote(doc,
          '结论：同一载荷、同一掩膜下，本程序与独立实现一致到峰值的'
          '万分之二到万分之三，可以放心用于定量分析。')

    # ---------------- 5
    h2(doc, '六、怎么获取')
    rich(doc, [
        ('安装包约 ', ''), ('52 MB', 'b'),
        ('，适用于 Windows 10/11（64 位），装完占约 200 MB 磁盘。'
         '建议内存 8 GB 以上。', ''),
    ])
    bullet(doc, '双击安装程序，按向导走完；桌面和开始菜单会出现快捷方式。')
    bullet(doc, '想知道装好没有，打开命令提示符运行：'
                '「安装目录\\sleqn.exe --selftest」，'
                '会检查随包文件、离线底图、数值内核与出图，'
                '最后打印「14/14 项通过」即一切就绪。')
    bullet(doc, '卸载走「设置 → 应用」，或用开始菜单里的卸载项。')

    quote(doc,
          '下载方式：【请在此处填写获取途径，例如"公众号后台回复「海平面指纹」'
          '获取安装包与示例数据"】')

    h2(doc, '七、两个使用提醒')
    bullet(doc, '命令行默认参数对真实 GRACE 数据会出错——默认网格间隔是 0.5°、'
                '且会把负值载荷清零。请写成 '
                '--dlon 1 --dlat 1 --allow-negative。'
                '（图形界面没有这个问题，界面默认值就是对的；'
                '程序现在也会在参数不匹配时打印警告。）',
           lead='注意符号与网格　')
    bullet(doc, '默认迭代 2 次与完全收敛解相差约 1%，够看形态；'
                '做趋势、区域平均这类定量分析时建议设到 10 次以上。',
           lead='要精度就多迭代　')

    # ---------------- 6
    h2(doc, '八、引用与反馈')
    rich(doc, [
        ('本程序实现并求解海平面方程，方法出自：', ''),
    ])
    rich(doc, [
        ('[1] 王林松, 陈超, 马险, 杜劲松. 冰盖消融的海平面指纹变化及其对 '
         'GRACE 监测结果的影响[J]. ', ''),
        ('地球物理学报', 'i'),
        (', 2018, 61(7): 2679–2690. doi:10.6038/cjg2018L0335', ''),
    ], size=10, space_after=6)
    rich(doc, [
        ('[2] Sun J. W., Wang L. S., Peng Z. R., Fu Z. Y., Chen C (2022). '
         'The sea level fingerprints of global terrestrial water storage changes '
         'detected by GRACE and GRACE-FO data. ', ''),
        ('Pure and Applied Geophysics', 'i'),
        (', 179(9), 3303–3317. doi:10.1007/s00024-022-03099-5', ''),
    ], size=10, space_after=12)
    rich(doc, [
        ('在论文或报告中使用了本程序的计算结果，请引用上述文献。', 'g'),
    ])
    rich(doc, [
        ('使用中遇到问题、发现异常结果，或者希望增加新功能，'
         '欢迎通过下面的邮箱或公众号留言反馈；反馈时附上程序「日志」页签的内容，'
         '便于快速定位。', ''),
    ])

    para(doc, '', space_after=4)
    image(doc, '', '扫码关注「地球重力与人类生活（TVGG）」，获取工具与更新',
          width=4.6, path=QR)

    rich(doc, [
        ('作者：彭桢燃（中国地质大学·武汉）　', 'g'),
        ('邮箱：zhenran.peng@cug.edu.cn', 'g'),
    ], size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    rich(doc, [
        ('程序完全离线运行，不联网、不上传任何数据。', 'g'),
    ], size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(OUT)
    return OUT


if __name__ == '__main__':
    p = build()
    n = len(open(p, 'rb').read())
    print(f'已生成 {os.path.relpath(p, ROOT)}  （{n / 1024:.0f} kB）')
    sys.exit(0)
